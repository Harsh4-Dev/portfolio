"""
upgrade_workbook.py — adds any Settings rows, Text rows, Sections/Journey columns that a newer
version of the site knows about to YOUR content.xlsx, without touching your content.

    python tools/upgrade_workbook.py            # upgrades ./content.xlsx in place (keeps a .bak copy)

Needs openpyxl (pip install openpyxl). Close the workbook in Excel first.
The site works without running this — every new key has a built-in default — this just makes the
new switches visible in the workbook so you can find and change them.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_seed import JOURNEY_HDR, SECTIONS_HDR, SETTINGS, TEXT, YESNO  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "content.xlsx"
HDR_FILL = PatternFill("solid", fgColor="2A2622")
HDR_FONT = Font(bold=True, color="F7EFE1")
COLUMN_DEFAULTS = {"themes": "all", "show_in_nav": "yes", "show_in_cv": "yes", "visible": "yes"}
COLUMN_OPTIONS = {"themes": ["all", "paper", "terminal"], "show_in_nav": YESNO, "show_in_cv": YESNO, "visible": YESNO}


def header_cell(ws, col, name):
    ws.cell(1, col).value = name
    ws.cell(1, col).fill = HDR_FILL
    ws.cell(1, col).font = HDR_FONT


def add_missing_columns(ws, wanted: list[str]) -> list[str]:
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    added = []
    for col in wanted:
        if col in hdr:
            continue
        c = len(hdr) + 1
        header_cell(ws, c, col)
        ws.column_dimensions[get_column_letter(c)].width = 22 if col == "command" else 12
        hdr.append(col)
        default = COLUMN_DEFAULTS.get(col, "")
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value not in (None, ""):
                ws.cell(r, c).value = default
        if col in COLUMN_OPTIONS:
            dv = DataValidation(type="list", formula1='"' + ",".join(COLUMN_OPTIONS[col]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(c)}2:{get_column_letter(c)}500")
        added.append(col)
    return added


def add_missing_rows(ws, rows: list[list], key_col=1) -> list[str]:
    have = {str(ws.cell(r, key_col).value).strip() for r in range(2, ws.max_row + 1)}
    added = []
    for row in rows:
        if str(row[0]) not in have:
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            added.append(row[0])
    return added


def main() -> int:
    if not XLSX.exists():
        print(f"{XLSX} not found")
        return 1
    bak = XLSX.with_suffix(".bak.xlsx")
    shutil.copy2(XLSX, bak)
    wb = load_workbook(XLSX)
    report = []
    if "Settings" in wb.sheetnames:
        report.append(("Settings rows", add_missing_rows(wb["Settings"], SETTINGS)))
    if "Text" not in wb.sheetnames:
        idx = wb.sheetnames.index("Links") + 1 if "Links" in wb.sheetnames else len(wb.sheetnames)
        tx = wb.create_sheet("Text", index=idx)
        tx.append(["key", "value", "note"])
        for c in range(1, 4):
            header_cell(tx, c, ["key", "value", "note"][c - 1])
        tx.column_dimensions["A"].width = 26
        tx.column_dimensions["B"].width = 90
        tx.column_dimensions["C"].width = 60
        tx.freeze_panes = "A2"
        tx.sheet_properties.tabColor = "6B6FD6"
        report.append(("Text sheet", ["created"]))
    report.append(("Text rows", add_missing_rows(wb["Text"], TEXT)))
    if "Sections" in wb.sheetnames:
        report.append(("Sections columns", add_missing_columns(wb["Sections"], SECTIONS_HDR)))
    if "Journey" in wb.sheetnames:
        report.append(("Journey columns", add_missing_columns(wb["Journey"], JOURNEY_HDR)))
    try:
        wb.save(XLSX)
    except PermissionError:
        print("content.xlsx is open in Excel - close it and run again (nothing was changed)")
        return 1
    for label, items in report:
        print(f"{label}: {', '.join(map(str, items)) if items else 'nothing to add'}")
    print(f"backup: {bak.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
