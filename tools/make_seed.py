"""
make_seed.py — (re)generates content.xlsx with the starter content.

You normally never need this: edit content.xlsx directly in Excel.
Run it only to reset the workbook to the seed (it OVERWRITES content.xlsx):

    python tools/make_seed.py            # writes ./content.xlsx
    python tools/make_seed.py out.xlsx   # writes somewhere else

Requires openpyxl (pip install openpyxl). The site builder itself needs nothing.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "content.xlsx"

INK = "2A2622"
CREAM = "F7EFE1"
CREAM2 = "EFE3CC"
HDR_FILL = PatternFill("solid", fgColor=INK)
HDR_FONT = Font(bold=True, color=CREAM, name="Calibri", size=11)
NOTE_FONT = Font(italic=True, color="7A7168", name="Calibri", size=10)
THIN = Side(style="thin", color="D9CDB6")
BORDER = Border(bottom=THIN)

LAYOUTS = ["timeline", "cards", "tags", "stats", "text", "list", "gallery", "table", "contact"]
SHAPES = ["neuron", "perceptron", "mlp", "cnn", "deep", "transformer", "constellation"]
YESNO = ["yes", "no"]
SHOW_IN = ["all", "nav", "hero", "footer", "contact", "none"]

# --------------------------------------------------------------------------- content

README_ROWS = [
    ["HOW THIS WORKBOOK WORKS", ""],
    ["", ""],
    ["Every sheet is a section of the site.", "Every row in it is one item (a job, a project, a skill...). Every column is a field."],
    ["Add a section", "Add a new sheet (its name becomes the section id) and, optionally, a row in `Sections` to pick its layout, title and order. Unregistered sheets still show up, as cards."],
    ["Add an item", "Add a row. Add a column if you need a new field - unknown columns are shown as small detail chips, so nothing is ever lost."],
    ["Remove something", "Delete the row, or put `yes` in a `hidden` column, or set `visible` = no in `Sections`."],
    ["Sheets starting with _", "are ignored by the site (like this one)."],
    ["", ""],
    ["RECOGNISED COLUMNS (all optional)", ""],
    ["title / subtitle", "Main and secondary line of an item."],
    ["org / role / location", "Organisation, role and place (timeline layouts show them next to the title)."],
    ["start / end / date", "Dates. Type `2026-06`, `Jun 2026`, `2026`, `Present`... anything - real Excel dates work too."],
    ["description", "Long text. Blank line = new paragraph, lines starting with `- ` = bullets, **bold**, *italic*, `code`, [link text](https://url)."],
    ["tags / tools", "Comma or semicolon separated. Rendered as doodle chips."],
    ["link / link_label", "One main link. `links` can hold several: `GitHub|https://... ; Paper|https://...`"],
    ["image / images", "File name inside site/assets/img/ (or a full URL). `images` = several, separated by ; "],
    ["icon", "An emoji or a doodle name (brain, chip, network, lightbulb, rocket, book, trophy, medal, mail, pin, coffee, star, sparkle, code, graph, heart, flask, globe)."],
    ["color", "Any CSS colour, or accent / accent2 / accent3 / accent4 to use the theme."],
    ["group", "For `tags` and `list` layouts: groups items under a small heading."],
    ["level", "1-5, for `tags` layout: draws filled dots (skill level)."],
    ["value / label", "For `stats` layout: the big number and its caption (value can be text like `20+`)."],
    ["featured", "yes = bigger card in `cards` layout."],
    ["order", "Number. Lower first. Otherwise the sheet order is kept."],
    ["hidden", "yes = not shown (keeps the row for later)."],
    ["", ""],
    ["LAYOUTS (Sections sheet)", ", ".join(LAYOUTS)],
    ["ERAS / SHAPES (Journey sheet)", ", ".join(SHAPES) + "  - each era is a moment in ML history; the 3D network morphs into that architecture when its section scrolls into view."],
    ["", ""],
    ["After editing", "Save. If `python studio.py` is running, the preview reloads by itself. Click Publish in the studio (or just git push) and GitHub builds the site."],
]

SETTINGS = [
    # key, value, note
    ["site_title", "Harshavardhan S - AI Systems Engineer", "Browser tab title"],
    ["name", "Harshavardhan S", "Full name (hero heading)"],
    ["first_name", "Harsh", "Used in casual spots (nav logo, footer)"],
    ["initials", "HS", "Nav logo"],
    ["tagline", "AI systems, end to end", "Short line under the name"],
    ["headline", "I fine-tune, serve and ship machine-learning systems - from domain-specific code LLMs on vLLM to a self-hosted vision app real people use every day.", "Hero paragraph"],
    ["roles", "domain-specific LLMs; self-hosted AI apps; clinical-trial data pipelines; full-stack ML tooling", "Rotating words in the hero, separated by ;"],
    ["roles_prefix", "I build", "Text before the rotating word"],
    ["availability", "Open to AI/ML internships & research roles", "Small pill above the name (leave blank to hide)"],
    ["location", "Trichy, Tamil Nadu, India", ""],
    ["email", "harsha.v27105@gmail.com", ""],
    ["phone", "", "Leave blank to keep it off the site"],
    ["avatar", "", "Image in site/assets/img/ e.g. me.jpg (optional)"],
    ["resume_url", "assets/Harshavardhan-S-Resume.pdf", "PDF path or URL for the CV button (blank = use cv.html)"],
    ["primary_cta_label", "See my work", ""],
    ["primary_cta_link", "#Projects", "#SectionId or URL"],
    ["secondary_cta_label", "Download CV", ""],
    ["secondary_cta_link", "cv.html", ""],
    ["hero_era", "perceptron", "Which Journey shape the 3D scene assembles into on load"],
    ["hero_note", "Every network on this page is a real architecture from ML history. Scroll - the story runs alongside mine.", "Small hand-written note near the 3D scene"],
    ["journey_title", "A short history of learning machines", "Title of the side rail"],
    ["show_journey_rail", "yes", "yes/no"],
    ["footer_note", "Drawn with cream, ink and gradient descent.", ""],
    ["seo_description", "Portfolio of Harshavardhan S - AI systems engineer: LLM fine-tuning, vLLM serving, self-hosted ML apps, data analytics.", ""],
    ["color_paper", "#f7efe1", "Toasty cream background"],
    ["color_paper_2", "#efe3cc", "Slightly darker cream (cards, rail)"],
    ["color_ink", "#2a2622", "Text & outlines"],
    ["color_accent", "#e2674a", "Coral"],
    ["color_accent_2", "#e8b43a", "Mustard"],
    ["color_accent_3", "#2f9e8f", "Teal"],
    ["color_accent_4", "#6b6fd6", "Periwinkle"],
    ["font_display", "Fraunces", "Any Google Font name"],
    ["font_hand", "Caveat", "Handwriting accents"],
    ["font_body", "Nunito", "Body text"],
    ["scene_density", "medium", "low / medium / high - number of nodes & pulses in the 3D scene"],
    ["scene_opacity_sections", "0.55", "How visible the 3D network stays behind the content (0-1)"],
    ["show_grain", "yes", "Paper grain texture"],
    ["cv_title", "Curriculum Vitae", "Heading on cv.html"],
    ["theme", "paper", "Which theme the site opens with: paper (cream, doodles, 3D) or terminal (dark CLI, ASCII). This is what GitHub Pages serves."],
    ["theme_toggle", "yes", "yes/no - let visitors switch between the two themes"],
    ["terminal_user", "harsh@portfolio", "Prompt shown in the terminal theme"],
    ["terminal_bg", "#000000", "Terminal background"],
    ["terminal_fg", "#ffffff", "Terminal text"],
    ["terminal_dim", "#8a8a8a", "Terminal muted text"],
    ["terminal_accent", "#f2c86b", "The one accent colour in the terminal theme (yellow): highlights, prompt caret, animation"],
    ["terminal_history", "yes", "yes/no - link to history.html (the ML-history page with the ASCII diagrams) from the terminal theme"],
    ["terminal_boot", "yes", "yes/no - short boot sequence on first load (terminal theme)"],
    ["hero_note_terminal", "", "Optional extra line under the hero in the terminal theme"],
    ["terminal_scanlines", "yes", "yes/no - CRT scanline overlay in the terminal theme"],
    ["terminal_font", "JetBrains Mono", "Any Google monospace font"],
]

SECTIONS_HDR = ["id", "title", "eyebrow", "intro", "layout", "order", "era", "doodle", "columns", "visible", "cta_label", "cta_link"]
SECTIONS = [
    ["About", "About me", "Hello", "", "text", 1, "neuron", "wave", "", "yes", "", ""],
    ["Numbers", "In numbers", "At a glance", "", "stats", 2, "neuron", "sparkle", "4", "yes", "", ""],
    ["Education", "Education", "Where I learned to learn", "", "timeline", 3, "perceptron", "book", "", "yes", "", ""],
    ["Skills", "Skills", "Toolbox", "Things I reach for most. Filled dots are how comfortable I am.", "tags", 4, "mlp", "chip", "", "yes", "", ""],
    ["Experience", "Experience", "Out in the field", "", "timeline", 5, "cnn", "pin", "", "yes", "", ""],
    ["Projects", "Projects", "Things I have built", "Click a card for the full story.", "cards", 6, "deep", "rocket", "2", "yes", "More on GitHub", "https://github.com/Harsh4-Dev"],
    ["Achievements", "Achievements", "Small trophies", "", "list", 7, "transformer", "trophy", "", "yes", "", ""],
    ["Certifications", "Certifications", "Paper trail", "", "list", 8, "transformer", "medal", "2", "yes", "", ""],
    ["Contact", "Let's build something", "Say hello", "I reply to every email. If you are working on LLM infrastructure, applied ML in healthcare, or anything self-hosted, I would love to hear about it.", "contact", 9, "constellation", "mail", "", "yes", "", ""],
]

JOURNEY_HDR = ["id", "year", "title", "insight", "shape", "color", "order"]
JOURNEY = [
    ["neuron", "1943", "The first artificial neuron",
     "Warren McCulloch and Walter Pitts show that a single threshold unit can compute logic. Every network on this page - and every model I have trained - is still built from that idea.",
     "neuron", "accent", 1],
    ["perceptron", "1958", "Rosenblatt's Mark I Perceptron",
     "A room-sized machine with a 20x20 grid of photocells and motor-driven potentiometers for weights. It learned to tell shapes apart from examples - the first machine that was taught rather than programmed.",
     "perceptron", "accent2", 2],
    ["mlp", "1986", "Backpropagation goes mainstream",
     "Rumelhart, Hinton and Williams show hidden layers can be trained by pushing errors backwards. Learning became a matter of gradients - which is still how every model I fine-tune improves.",
     "mlp", "accent3", 3],
    ["cnn", "1998", "LeNet-5 reads the cheques",
     "Yann LeCun's convolutional network shipped inside bank machines and at one point read a sizeable share of US cheques. The first neural net doing quiet, industrial-scale work in production.",
     "cnn", "accent4", 4],
    ["deep", "2012", "AlexNet and the GPU era",
     "Two GTX 580 cards, a network split in half across them, and the ImageNet error rate drops by a third. GPUs became the substrate of the field - and, later, of my own self-hosted stack.",
     "deep", "accent", 5],
    ["transformer", "2017", "Attention is all you need",
     "Vaswani et al. replace recurrence with attention: every token looks at every other token. The architecture behind the LLMs I fine-tune and serve today.",
     "transformer", "accent3", 6],
    ["constellation", "Now", "Your chapter",
     "Models are cheap to run and expensive to get right. The interesting work is the whole system: data, weights, serving, the people using it.",
     "constellation", "accent2", 7],
]

LINKS_HDR = ["label", "url", "icon", "show_in", "order"]
LINKS = [
    ["GitHub", "https://github.com/Harsh4-Dev", "github", "all", 1],
    ["LinkedIn", "https://www.linkedin.com/in/harshav1729", "linkedin", "all", 2],
    ["Email", "mailto:harsha.v27105@gmail.com", "mail", "all", 3],
]

ABOUT_HDR = ["title", "description", "image", "order"]
ABOUT = [
    ["Hi, I'm Harsh.",
     "I'm a Computer Science undergraduate (AI & Data Science) at SASTRA and a Data Analytics intern at Roche, where I work with clinical-trial data.\n\n"
     "What I enjoy most is the full arc of an ML system: getting data into shape, fine-tuning a model with PEFT/DoRA, serving it fast on vLLM, and wrapping it in something people actually use. My nutrition tracker runs on a GPU box I administer myself and has 20+ daily users who never think about the model behind the camera.\n\n"
     "Outside of work I read ML history (hence this page), tinker with self-hosting, and keep a long list of things to rebuild from scratch just to understand them.",
     "", 1],
]

NUMBERS_HDR = ["value", "label", "icon", "order"]
NUMBERS = [
    ["20+", "daily users on my self-hosted AI app", "heart", 1],
    ["4", "LLM training paradigms used (FIM, CLM, SFT, DoRA)", "brain", 2],
    ["Top 2%", "of 33,000+ in NPTEL Cloud Computing", "trophy", 3],
    ["8.25", "CGPA, B.Tech CSE (AI & DS)", "book", 4],
]

TIMELINE_HDR = ["title", "org", "location", "start", "end", "description", "tags", "link", "icon", "order"]
EXPERIENCE = [
    ["Data Analytics Intern", "Roche", "Remote / India", "2026-06", "Present",
     "- Query, clean and analyse large clinical-trial datasets with SQL and Python (pandas, NumPy) to produce reliable, analysis-ready data.\n"
     "- Surface trends and anomalies through exploratory analysis and ML-assisted pattern detection to inform study-team decisions.\n"
     "- Automate recurring reporting and data-quality checks, improving consistency and turnaround for downstream stakeholders.",
     "SQL, Python, pandas, NumPy, clinical trials", "", "flask", 1],
]
EDUCATION = [
    ["B.Tech - Computer Science & Engineering (AI & Data Science)", "SASTRA Deemed to be University", "Thanjavur, India", "2023", "2027",
     "CGPA 8.25. Specialisation in artificial intelligence and data science.", "", "", "book", 1],
    ["Honors - Applied Accelerated Artificial Intelligence", "NPTEL - IIT Guwahati", "", "2025", "",
     "GPU compute and system software for accelerated-computing AI solutions, with end-to-end industrial deployment. Elite certification.", "GPU, CUDA, deployment", "", "chip", 2],
    ["Honors - Cloud Computing", "NPTEL - IIT Kharagpur", "", "2025", "",
     "Cloud architectures, virtualisation and service models. Top 2% of 33,000+ certified candidates (Elite - Topper, 90%).", "cloud, virtualisation", "", "globe", 3],
]

SKILLS_HDR = ["title", "group", "level", "order"]
SKILLS = [
    ["Python", "Languages", 5, 1], ["SQL", "Languages", 4, 2], ["JavaScript", "Languages", 4, 3], ["TypeScript", "Languages", 3, 4],
    ["PyTorch", "AI / ML", 4, 1], ["Hugging Face", "AI / ML", 5, 2], ["PEFT / DoRA", "AI / ML", 4, 3], ["vLLM", "AI / ML", 4, 4],
    ["LangChain", "AI / ML", 3, 5], ["Computer Vision", "AI / ML", 3, 6], ["pandas / NumPy", "AI / ML", 4, 7], ["RAG & document parsing", "AI / ML", 4, 8],
    ["FastAPI", "Frameworks", 5, 1], ["Django", "Frameworks", 3, 2], ["React", "Frameworks", 3, 3], ["REST APIs", "Frameworks", 4, 4],
    ["Docker", "Infra & tools", 4, 1], ["Linux", "Infra & tools", 5, 2], ["nginx", "Infra & tools", 3, 3], ["Git", "Infra & tools", 4, 4],
    ["Cloudflare Tunnel", "Infra & tools", 3, 5], ["SQLite", "Infra & tools", 4, 6], ["NVIDIA DGX", "Infra & tools", 3, 7],
]

PROJECTS_HDR = ["title", "subtitle", "description", "tags", "link", "link_label", "links", "image", "featured", "icon", "color", "date", "status", "order", "hidden"]
PROJECTS = [
    ["Domain-Specific Code LLM", "Fine-tuned code model for a proprietary enterprise platform",
     "An end-to-end pipeline to build a code LLM for a proprietary enterprise development language.\n\n"
     "- Trained base models across four paradigms - fill-in-the-middle (FIM), causal LM, supervised fine-tuning and DoRA - and served them at high throughput with **vLLM** behind an OpenWebUI chat interface.\n"
     "- Rebuilt the platform's core from scratch: a functional re-implementation of its multi-value database engine (maker-checker approval lifecycle, write-ahead-log crash recovery) exposed through a REST API, used to synthesise and validate training data.\n"
     "- Built an autonomous full-stack coding agent that converts legacy code into the target language, expanding the training corpus and closing the fine-tuning loop.",
     "Python, PyTorch, Hugging Face, PEFT/DoRA, vLLM, OpenWebUI, FastAPI, Linux", "", "", "", "", "yes", "brain", "accent", "2025 - 2026", "Private", 1, ""],
    ["Harahachi", "Self-hosted AI calorie & macro tracker - 20+ active users",
     "A photo-first nutrition tracker: a vision LLM identifies each food item on the plate and estimates calories and macros; you confirm with one tap.\n\n"
     "- **FastAPI edge-AI gateway** with QR device pairing, multi-user accounts with per-user isolation, a pluggable multi-LLM registry (vLLM / LM Studio / NVIDIA API) with automatic failover, push reminders and an admin dashboard.\n"
     "- Packaged as a single Docker container (nginx + FastAPI + outbound-only Cloudflare tunnel), deployed rootless on an NVIDIA DGX.\n"
     "- Shipped as an installable PWA and a Capacitor Android app, backed by a 150+ test suite.",
     "Python, FastAPI, SQLite, Docker, nginx, Cloudflare Tunnel, JavaScript, Capacitor, NVIDIA NIM", "https://github.com/Harsh4-Dev/harahachi", "GitHub", "", "", "yes", "heart", "accent3", "2025 - now", "In production", 2, ""],
    ["docparser", "Any document to RAG-ready chunks, offline, on CPU",
     "Turns PDFs, scans, images, Office files and HTML into retrieval-ready chunks with tables, figures, fonts and colour metadata attached.\n\n"
     "- Per-page routing between native text extraction and OCR (RapidOCR + Tesseract consensus that re-inserts lost word boundaries), with the decision and its evidence recorded per page.\n"
     "- Four table detectors (ruled, borderless, pixel grids, format-native) normalised onto one representation, with cross-page merging and linearised rows so table content is embeddable.\n"
     "- Structure-driven chunking - never a sliding window - under a ~1 GB RAM ceiling.",
     "Python, PyMuPDF, ONNX, RapidOCR, Tesseract, RAG", "", "", "", "", "", "network", "accent4", "2026", "Active", 3, ""],
    ["PORTOLAN", "Full-stack agentic LLM factory",
     "Describe PORTOLAN here - it is hidden until you fill this in and clear the `hidden` column.",
     "Python, agents", "", "", "", "", "", "rocket", "accent2", "2026", "WIP", 4, "yes"],
]

LIST_HDR = ["title", "subtitle", "date", "description", "link", "icon", "order"]
ACHIEVEMENTS = [
    ["Finalist - ANRF AISEHack 2026", "Anusandhan National Research Foundation with IBM Research & IIT Delhi, held at IIIT Hyderabad", "2026",
     "Team Hackaholics.", "", "trophy", 1],
    ["Elite - Topper, NPTEL Cloud Computing", "Top 2% of 33,000+ certified candidates, score 90%", "2025", "", "", "medal", 2],
]
CERTIFICATIONS = [
    ["Hugging Face LLM Course", "Fundamentals of LLMs; Fine-tuning Language Models", "2025", "", "https://huggingface.co/learn", "brain", 1],
    ["Applied Accelerated Artificial Intelligence", "NPTEL - IIT Guwahati, Elite", "2025", "", "", "chip", 2],
    ["Machine Learning & Deep Learning Onramp", "MathWorks", "2024", "", "", "graph", 3],
    ["Graph Theory Programming Camp & DSA Course", "AlgoUniversity, mentored by a Codeforces Master", "2024", "", "", "code", 4],
]

CONTACT_HDR = ["title", "subtitle", "link", "icon", "order"]
CONTACT = [
    ["Email", "harsha.v27105@gmail.com", "mailto:harsha.v27105@gmail.com", "mail", 1],
    ["GitHub", "github.com/Harsh4-Dev", "https://github.com/Harsh4-Dev", "github", 2],
    ["LinkedIn", "linkedin.com/in/harshav1729", "https://www.linkedin.com/in/harshav1729", "linkedin", 3],
]

# --------------------------------------------------------------------------- writer helpers


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def autosize(ws, widths: dict[str, int] | None = None, default=18, wrap_cols=()):
    widths = widths or {}
    hdr = [c.value for c in ws[1]]
    for i, h in enumerate(hdr, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths.get(h, default)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            h = hdr[cell.column - 1] if cell.column - 1 < len(hdr) else None
            cell.alignment = Alignment(wrap_text=(h in wrap_cols), vertical="top")
            cell.border = BORDER


def add_sheet(wb, name, header, rows, widths=None, wrap=("description", "intro", "insight", "note", "headline", "value")):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    style_header(ws, len(header))
    autosize(ws, widths, wrap_cols=wrap)
    ws.sheet_properties.tabColor = CREAM2
    return ws


def add_validation(ws, header, col_name, options, max_row=500):
    if col_name not in header:
        return
    col = get_column_letter(header.index(col_name) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showDropDown=False)
    dv.error = f"Pick one of: {', '.join(options)}"
    dv.errorTitle = "Not a valid value"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{max_row}")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "_README"
    for r in README_ROWS:
        ws.append(r)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=INK)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].font = NOTE_FONT
    for r in (9, 26):
        ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="E2674A")
    ws.sheet_properties.tabColor = "E2674A"

    s = add_sheet(wb, "Settings", ["key", "value", "note"], SETTINGS, {"key": 26, "value": 70, "note": 60}, wrap=("value", "note"))
    s.sheet_properties.tabColor = "6B6FD6"

    sec = add_sheet(wb, "Sections", SECTIONS_HDR, SECTIONS, {"id": 16, "title": 24, "eyebrow": 22, "intro": 50, "layout": 12, "order": 7, "era": 14, "doodle": 10, "columns": 9, "visible": 8, "cta_label": 16, "cta_link": 34})
    sec.sheet_properties.tabColor = "6B6FD6"
    add_validation(sec, SECTIONS_HDR, "layout", LAYOUTS)
    add_validation(sec, SECTIONS_HDR, "era", [j[0] for j in JOURNEY])
    add_validation(sec, SECTIONS_HDR, "visible", YESNO)

    jr = add_sheet(wb, "Journey", JOURNEY_HDR, JOURNEY, {"id": 14, "year": 8, "title": 34, "insight": 90, "shape": 14, "color": 10, "order": 7})
    jr.sheet_properties.tabColor = "6B6FD6"
    add_validation(jr, JOURNEY_HDR, "shape", SHAPES)

    ln = add_sheet(wb, "Links", LINKS_HDR, LINKS, {"label": 14, "url": 48, "icon": 10, "show_in": 10, "order": 7})
    ln.sheet_properties.tabColor = "6B6FD6"
    add_validation(ln, LINKS_HDR, "show_in", SHOW_IN)

    add_sheet(wb, "About", ABOUT_HDR, ABOUT, {"title": 22, "description": 100, "image": 16, "order": 7})
    add_sheet(wb, "Numbers", NUMBERS_HDR, NUMBERS, {"value": 10, "label": 44, "icon": 10, "order": 7})
    tl_w = {"title": 40, "org": 30, "location": 18, "start": 10, "end": 10, "description": 90, "tags": 30, "link": 30, "icon": 8, "order": 7}
    add_sheet(wb, "Education", TIMELINE_HDR, EDUCATION, tl_w)
    add_sheet(wb, "Skills", SKILLS_HDR, SKILLS, {"title": 26, "group": 16, "level": 7, "order": 7})
    add_sheet(wb, "Experience", TIMELINE_HDR, EXPERIENCE, tl_w)
    pj = add_sheet(wb, "Projects", PROJECTS_HDR, PROJECTS, {"title": 28, "subtitle": 44, "description": 100, "tags": 40, "link": 34, "link_label": 10, "links": 30, "image": 16, "featured": 9, "icon": 8, "color": 9, "date": 12, "status": 12, "order": 7, "hidden": 8})
    add_validation(pj, PROJECTS_HDR, "featured", YESNO)
    add_validation(pj, PROJECTS_HDR, "hidden", YESNO)
    add_sheet(wb, "Achievements", LIST_HDR, ACHIEVEMENTS, {"title": 40, "subtitle": 60, "date": 8, "description": 50, "link": 30, "icon": 8, "order": 7})
    add_sheet(wb, "Certifications", LIST_HDR, CERTIFICATIONS, {"title": 44, "subtitle": 50, "date": 8, "description": 50, "link": 30, "icon": 8, "order": 7})
    add_sheet(wb, "Contact", CONTACT_HDR, CONTACT, {"title": 14, "subtitle": 36, "link": 44, "icon": 8, "order": 7})

    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
